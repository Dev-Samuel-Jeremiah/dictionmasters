"""A school admin moving and promoting their teachers and students."""

from django.test import TestCase

from apps.accounts.models import User

from .models import LevelChange, School


class LevelToolsTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name="Unity School", email="unity@example.com")
        self.admin = User.objects.create_user("admin@example.com", "pw-12345678", first_name="Ada",
                                              role="school_admin", school=self.school)
        self.client.force_login(self.admin)

    def person(self, name, level="", role="student", school=None, additional=""):
        return User.objects.create_user(f"{name}@example.com", "pw-12345678", first_name=name.title(), role=role,
                                        school=school or self.school, level=level, additional_levels=additional)

    def bulk(self, people, action, level="", role="student"):
        return self.client.post("/school/members/levels/", {"members": [p.pk for p in people], "action": action,
                                                             "level": level, "role": role}, follow=True)

    def levels(self, *people):
        return [User.objects.get(pk=p.pk).level for p in people]

    def test_promote_moves_each_student_up_one_level(self):
        pre, three, top, none = (self.person("pre", "Pre-Level"), self.person("three", "Level 3"),
                                 self.person("top", "Level 12"), self.person("none"))
        response = self.bulk([pre, three, top, none], "promote")
        self.assertEqual(self.levels(pre, three, top, none), ["Level 1", "Level 4", "Level 12", ""])
        self.assertContains(response, "already in Level 12 stayed there")
        self.assertContains(response, "with no level yet was left")
        self.assertEqual(LevelChange.objects.count(), 2)

    def test_undo_puts_a_whole_promotion_back_except_people_moved_since(self):
        a, b = self.person("a", "Level 2"), self.person("b", "Level 5")
        self.bulk([a, b], "promote")
        batch = LevelChange.objects.first().batch
        self.client.post(f"/school/members/{b.pk}/level/", {"levels": ["Level 9"]})     # changed again
        response = self.client.post(f"/school/levels/undo/{batch}/", follow=True)
        self.assertEqual(self.levels(a, b), ["Level 2", "Level 9"])
        self.assertContains(response, "had been moved again since")
        self.assertTrue(LevelChange.objects.filter(batch=batch, undone_at__isnull=False).exists())

    def test_move_teachers_keeps_their_other_levels(self):
        teacher = self.person("tee", "Level 1", role="teacher", additional=",Level 3,Level 6,")
        self.bulk([teacher], "move", "Level 3", role="teacher")
        teacher.refresh_from_db()
        self.assertEqual((teacher.level, teacher.additional_levels_list), ("Level 3", ["Level 6"]))

    def test_a_teacher_can_be_given_several_levels(self):
        teacher = self.person("tee", "Level 4", role="teacher")
        self.client.post(f"/school/members/{teacher.pk}/level/", {"levels": ["Level 2", "Level 4", "Level 7"]})
        teacher.refresh_from_db()
        self.assertEqual(teacher.level, "Level 4")                      # still their main level
        self.assertEqual(teacher.all_levels, ["Level 2", "Level 4", "Level 7"])

    def test_a_student_has_one_level(self):
        student = self.person("stu", "Level 4")
        self.client.post(f"/school/members/{student.pk}/level/", {"levels": ["Level 6", "Level 8"]})
        student.refresh_from_db()
        self.assertEqual((student.level, student.additional_levels), ("Level 6", ""))

    def test_teachers_are_not_promoted_and_other_schools_are_untouched(self):
        other_school = School.objects.create(name="Elsewhere", email="else@example.com")
        outsider = self.person("out", "Level 2", school=other_school)
        teacher = self.person("tee", "Level 2", role="teacher")
        self.bulk([outsider], "promote")
        self.bulk([teacher], "promote", role="teacher")
        self.assertEqual(self.levels(outsider, teacher), ["Level 2", "Level 2"])

    def test_an_unknown_level_is_refused(self):
        student = self.person("stu", "Level 4")
        self.bulk([student], "move", "Level 99")
        self.assertEqual(self.levels(student), ["Level 4"])

    def test_only_school_admins_can_change_levels(self):
        student = self.person("stu", "Level 4")
        teacher = self.person("tee", "Level 4", role="teacher")
        self.client.force_login(teacher)
        self.bulk([student], "promote")
        self.assertEqual(self.levels(student), ["Level 4"])

    def test_the_dashboard_shows_the_tools_and_history(self):
        student = self.person("stu", "Level 4")
        self.bulk([student], "promote")
        response = self.client.get("/school/dashboard/")
        self.assertContains(response, "Promote to next level")
        self.assertContains(response, "Level 4 &rarr; Level 5")
        self.assertContains(response, "Undo")


import io

from cryptography.fernet import Fernet
from django.test import override_settings
from openpyxl import load_workbook

from apps.accounts import credentials

KEY = Fernet.generate_key().decode()


@override_settings(SCHOOL_CREDENTIALS_ENCRYPTION_KEY=KEY)
class LoginToolsTests(TestCase):
    def setUp(self):
        credentials._fernet.cache_clear()
        self.school = School.objects.create(name="Unity School", email="unity@example.com")
        self.admin = User.objects.create_user("admin@example.com", "admin-pass-123", first_name="Ada",
                                              role="school_admin", school=self.school)
        self.teacher = User.objects.create_user("t@example.com", "Teach-pass-99", first_name="Tola", role="teacher",
                                                school=self.school, level="Level 3")
        self.student = User.objects.create_user("s@example.com", "mango47", first_name="Sade", role="student",
                                                school=self.school, level="Level 3", username="sade.k")
        self.client.force_login(self.admin)

    def tearDown(self):
        credentials._fernet.cache_clear()

    def unlock(self, password="admin-pass-123"):
        return self.client.post("/school/logins/unlock/", {"password": password}, follow=True)

    def sheet(self, response):
        rows = load_workbook(io.BytesIO(response.content)).active.iter_rows(min_row=6, values_only=True)
        return {row[5]: row[6] for row in rows if row[5]}

    def test_everything_is_locked_until_the_admin_confirms_their_password(self):
        ajax = {"HTTP_X_REQUESTED_WITH": "XMLHttpRequest"}
        self.assertEqual(self.client.get(f"/school/members/{self.student.pk}/login/", **ajax).status_code, 403)
        self.assertEqual(self.client.post(f"/school/members/{self.student.pk}/reset-password/", **ajax).status_code, 403)
        self.assertRedirects(self.client.post("/school/logins/download/"), "/school/dashboard/#logins",
                             fetch_redirect_response=False)
        self.assertContains(self.unlock("wrong"), "isn&#x27;t right")
        self.assertContains(self.unlock(), "unlocked for 15 minutes")
        self.assertEqual(self.client.get(f"/school/members/{self.student.pk}/login/").json()["password"], "mango47")

    def test_reset_makes_a_new_simple_password_and_signs_them_out(self):
        self.unlock()
        self.student.device_logins.create(key_hash="x" * 64, auth_hash="y")
        data = self.client.post(f"/school/members/{self.student.pk}/reset-password/").json()
        self.student.refresh_from_db()
        self.assertTrue(self.student.check_password(data["password"]))
        self.assertFalse(self.student.check_password("mango47"))
        self.assertFalse(self.student.device_logins.exists())
        self.assertEqual(data["login"], "sade.k")

    def test_an_admin_can_choose_the_new_password(self):
        self.unlock()
        self.client.post(f"/school/members/{self.student.pk}/reset-password/", {"password": "banana22"})
        self.student.refresh_from_db()
        self.assertTrue(self.student.check_password("banana22"))
        # Simple passwords are fine now (apps/accounts/password_rules.py); too short isn't.
        short = self.client.post(f"/school/members/{self.teacher.pk}/reset-password/", {"password": "12345"})
        self.assertEqual(short.status_code, 400)

    def test_download_everyone_or_just_students_as_excel(self):
        self.unlock()
        everyone = self.sheet(self.client.post("/school/logins/download/"))
        self.assertEqual(everyone, {"t@example.com": "Teach-pass-99", "sade.k": "mango47"})
        self.assertNotIn("admin@example.com", everyone)
        students = self.sheet(self.client.post("/school/logins/download/", {"role": "student", "level": "Level 3"}))
        self.assertEqual(list(students), ["sade.k"])

    def test_passwords_that_cant_be_shown_can_be_filled_in(self):
        User.objects.filter(pk=self.student.pk).update(encrypted_login_password="")
        self.unlock()
        before = self.sheet(self.client.post("/school/logins/download/", {"role": "student"}))
        self.assertTrue(before["sade.k"].startswith("Not available"))
        after = self.sheet(self.client.post("/school/logins/download/", {"role": "student", "fill_missing": "1"}))
        self.student.refresh_from_db()
        self.assertTrue(self.student.check_password(after["sade.k"]))

    def test_resetting_the_ticked_people_downloads_their_new_passwords(self):
        self.unlock()
        response = self.client.post("/school/logins/reset/", {"members": [self.student.pk, self.teacher.pk]})
        self.assertIn("new-passwords", response["Content-Disposition"])
        new = self.sheet(response)
        self.teacher.refresh_from_db()
        self.assertTrue(self.teacher.check_password(new["t@example.com"]))
        self.assertFalse(self.teacher.check_password("Teach-pass-99"))

    def test_other_schools_and_non_admins_are_refused(self):
        other = User.objects.create_user("o@example.com", "pw-12345678", first_name="Oti", role="student",
                                         school=School.objects.create(name="Elsewhere", email="e@example.com"))
        self.unlock()
        self.assertEqual(self.client.get(f"/school/members/{other.pk}/login/").status_code, 404)
        self.assertEqual(self.client.post("/school/logins/reset/", {"members": [other.pk]}).status_code, 302)
        other.refresh_from_db()
        self.assertTrue(other.check_password("pw-12345678"))           # untouched
        self.client.force_login(self.teacher)
        self.assertNotEqual(self.client.get(f"/school/members/{self.student.pk}/login/").status_code, 200)
