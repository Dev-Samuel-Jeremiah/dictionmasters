"""Phase B of classes: the school admin's Classes pages, students joining a
class when they sign up, bulk add by grade and class, and the control
room's Classes screen."""

import io

from django.contrib.auth import get_user_model
from django.test import TestCase

from apps.manage import bulk_students
from apps.scheme.models import Session

from .models import LevelChange, School, SchoolClass

User = get_user_model()


def person(email, **extra):
    return User.objects.create_user(email, "mango-river-47", first_name=email.split("@")[0].title(), **extra)


class ClassCase(TestCase):
    def setUp(self):
        self.session = Session.objects.create(name="2026/2027")
        self.school = School.objects.create(name="Unity", email="u@example.com")
        self.admin = person("head@example.com", role="school_admin", school=self.school)
        self.teacher = person("tola@example.com", role="teacher", school=self.school, level="Level 3")
        self.ada = person("ada@example.com", role="student", school=self.school, level="Level 3")
        self.p3 = self.ada.school_class                       # "Grade 3", made for her


class ClassesPageTests(ClassCase):
    def setUp(self):
        super().setUp()
        self.client.force_login(self.admin)

    def test_the_classes_are_listed_by_stage_with_their_links(self):
        SchoolClass.objects.create(school=self.school, session=self.session, grade="nursery_1")
        page = self.client.get("/school/classes/")
        self.assertContains(page, "Pre-School")
        self.assertContains(page, "Grade 3")
        self.assertContains(page, f"/accounts/register/student/?school={self.school.code}&amp;class={self.p3.pk}")
        self.assertLess(page.content.index(b"Nursery 1"), page.content.index(b"Grade 3"))
        self.assertContains(self.client.get("/school/dashboard/"), 'href="/school/classes/"')

    def test_adding_a_class_with_its_teacher(self):
        response = self.client.post("/school/classes/", {"grade": "primary_3", "name": "  Gold ", "class_teacher": self.teacher.pk})
        gold = SchoolClass.objects.get(name="Gold")
        self.assertRedirects(response, f"/school/classes/{gold.pk}/")
        self.assertEqual((gold.school, gold.session, gold.display_name), (self.school, self.session, "Grade 3 Gold"))
        self.assertEqual(list(gold.teachers.all()), [self.teacher])     # the class teacher teaches it too
        again = self.client.post("/school/classes/", {"grade": "primary_3", "name": "gold"})
        self.assertContains(again, "There&#x27;s already a Grade 3 gold class this year")

    def test_only_the_schools_own_teachers_can_be_chosen(self):
        stranger = person("x@example.com", role="teacher", school=School.objects.create(name="Hope", email="h@example.com"))
        page = self.client.post("/school/classes/", {"grade": "primary_4", "class_teacher": stranger.pk})
        self.assertFalse(SchoolClass.objects.filter(grade="primary_4").exists())
        self.assertContains(page, "Select a valid choice")

    def test_moving_students_to_another_class_changes_their_level(self):
        jss = SchoolClass.objects.create(school=self.school, session=self.session, grade="jss_1", name="A")
        self.client.post(f"/school/classes/{self.p3.pk}/", {"action": "move", "students": [self.ada.pk], "to": jss.pk})
        self.ada.refresh_from_db()
        self.assertEqual((self.ada.school_class, self.ada.level), (jss, "Level 7"))
        change = LevelChange.objects.get(member=self.ada)
        self.assertEqual((change.old_level, change.new_level, change.kind), ("Level 3", "Level 7", "move"))

    def test_renaming_and_deleting(self):
        self.client.post(f"/school/classes/{self.p3.pk}/", {"action": "save", "grade": "primary_3", "name": "Blue"})
        self.p3.refresh_from_db()
        self.assertEqual(self.p3.display_name, "Grade 3 Blue")
        kept = self.client.post(f"/school/classes/{self.p3.pk}/", {"action": "delete"}, follow=True)
        self.assertContains(kept, "still has students")
        empty = SchoolClass.objects.create(school=self.school, session=self.session, grade="sss_1")
        self.assertRedirects(self.client.post(f"/school/classes/{empty.pk}/", {"action": "delete"}), "/school/classes/")
        self.assertFalse(SchoolClass.objects.filter(pk=empty.pk).exists())

    def test_only_this_schools_admin(self):
        other = SchoolClass.objects.create(school=School.objects.create(name="Hope", email="h@example.com"),
                                           grade="primary_1")
        self.assertEqual(self.client.get(f"/school/classes/{other.pk}/").status_code, 404)
        self.client.force_login(self.teacher)
        self.assertEqual(self.client.get("/school/classes/").status_code, 302)


class JoiningAClassTests(ClassCase):
    url = "/accounts/register/student/"

    def form(self, choice, **extra):
        return {"school_code": self.school.code, "school_class": choice, "first_name": "Bo", "username": "bo.okafor",
                "password1": "kola-river-9", "password2": "kola-river-9", **extra}

    def test_the_classes_link_chooses_the_class(self):
        from urllib.parse import quote

        from apps.accounts.views import TOUR_DONE

        link = f"{self.url}?school={self.school.code}&class={self.p3.pk}"
        # A first visit goes through the welcome tour, which brings them back to the link.
        self.assertIn(quote(link, safe=""), self.client.get(link)["Location"])
        session = self.client.session
        session[TOUR_DONE] = True
        session.save()
        page = self.client.get(link)
        self.assertContains(page, f'<option value="class:{self.p3.pk}" selected>Grade 3</option>', html=True)
        self.assertContains(page, '<option value="grade:jss_1">Grade 7</option>', html=True)    # no class yet

    def test_joining_a_class(self):
        self.client.post(self.url, self.form(f"class:{self.p3.pk}"))
        bo = User.objects.get(username="bo.okafor")
        self.assertEqual((bo.school, bo.school_class, bo.level), (self.school, self.p3, "Level 3"))

    def test_joining_a_grade_with_no_class_yet_makes_it(self):
        self.client.post(self.url, self.form("grade:nursery_2"))
        bo = User.objects.get(username="bo.okafor")
        self.assertEqual((bo.school_class.display_name, bo.level), ("Nursery 2", "Pre-Level"))

    def test_another_schools_class_is_refused(self):
        theirs = SchoolClass.objects.create(school=School.objects.create(name="Hope", email="h@example.com"), grade="primary_2")
        page = self.client.post(self.url, self.form(f"class:{theirs.pk}"))
        self.assertFalse(User.objects.filter(username="bo.okafor").exists())
        self.assertEqual(page.status_code, 200)

    def test_the_classes_for_a_typed_code(self):
        found = self.client.get("/accounts/register/student/classes/", {"school": self.school.code.removeprefix("DM-").lower()})
        groups = {g["label"]: g["options"] for g in found.json()["groups"]}
        self.assertIn([f"class:{self.p3.pk}", "Grade 3"], groups["Primary"])
        self.assertIn(["grade:creche", "Creche"], groups["Pre-School"])
        self.assertEqual(self.client.get("/accounts/register/student/classes/", {"school": "DM-NOPE"}).status_code, 404)


class BulkAddTests(ClassCase):
    def test_grades_and_class_names_in_the_file(self):
        csv = ("First name,Last name,Class (grade),Email,Class name\n"
               "Ada,Obi,Primary 3,,\nBo,Ali,JSS1,,Gold\nCy,Eze,Level 4,,\nDi,Umar,2,,\nEd,Bad,Primary 9,,\n").encode()
        upload = io.BytesIO(csv)
        upload.name = "class.csv"
        clean, problems = bulk_students.check(self.school, bulk_students.read_rows(upload))
        self.assertEqual([p[0] for p in problems], [6])                   # Primary 9 isn't a grade
        self.assertEqual([(r["level"], r["grade"], r["class_name"]) for r in clean],
                         [("Level 3", "primary_3", ""), ("Level 7", "jss_1", "Gold"),
                          ("Level 4", "primary_4", ""), ("Level 2", "primary_2", "")])
        bulk_students.create(self.school, clean)
        placed = {u.first_name: u.school_class.display_name for u in User.objects.filter(school=self.school, role="student")
                  .select_related("school_class")}
        self.assertEqual(placed, {"Ada": "Grade 3", "Bo": "Grade 7 Gold", "Cy": "Grade 4", "Di": "Grade 2"})
        self.assertEqual(User.objects.get(first_name="Ada", last_name="Obi").school_class, self.p3)

    def test_the_student_template_lists_grades_and_has_a_class_name_column(self):
        from openpyxl import load_workbook

        book = load_workbook(io.BytesIO(bulk_students.template_file(self.school)))
        headings = [c.value for c in book["Students"][bulk_students.FIRST_DATA_ROW - 1]]
        self.assertIn("Class (grade) *", headings)
        self.assertIn("Class name (optional)", headings)
        self.assertEqual(book["Levels"]["A1"].value, "Creche")
        teachers = load_workbook(io.BytesIO(bulk_students.template_file(self.school, "teacher")))
        self.assertEqual(teachers["Levels"]["A1"].value, "Pre-Level")


class ControlRoomTests(ClassCase):
    def test_staff_see_a_schools_classes(self):
        self.client.force_login(person("staff@example.com", is_staff=True))
        self.assertContains(self.client.get("/manage/classes/"), "Grade 3")
        self.assertEqual(self.client.get("/manage/classes/new/").status_code, 200)
        self.assertContains(self.client.get(f"/manage/schools/{self.school.pk}/"), "Classes")

    def test_the_classes_and_grades_page(self):
        self.client.force_login(person("staff@example.com", is_staff=True))
        page = self.client.get("/manage/classes-grades/")
        self.assertContains(page, "<td>Grade 7</td>", html=True)                  # the grades and their levels
        self.assertContains(page, "<td>Level 7</td>", html=True)
        self.assertContains(page, "Unity")
        self.assertContains(page, f'href="/manage/classes/{self.p3.pk}/"')       # each class to edit
        self.assertContains(page, 'href="/manage/classes-grades/"')              # in the sidebar
        other = School.objects.create(name="Hope", email="h@example.com")
        SchoolClass.objects.create(school=other, grade="sss_2", name="Lion")
        one = self.client.get(f"/manage/classes-grades/?school={other.pk}")
        self.assertContains(one, "Grade 11 Lion")
        self.assertNotContains(one, f'href="/manage/classes/{self.p3.pk}/"')
        self.assertContains(one, f'href="/manage/classes/new/?in={other.pk}"')
        self.assertContains(self.client.get("/manage/"), 'href="/manage/classes/new/"')   # Add something new
        self.assertContains(self.client.get(f"/manage/schools-directory/{self.school.pk}/"),
                            f'href="/manage/classes-grades/?school={self.school.pk}"')


class GradeEditTests(ClassCase):
    """Staff renaming a grade, or changing the level it follows."""

    def setUp(self):
        super().setUp()
        from . import grades

        self.addCleanup(grades.forget)              # the grade cache outlives a test's database
        self.client.force_login(person("staff@example.com", is_staff=True))
        self.p3.teachers.add(self.teacher)

    def test_every_grade_has_an_edit_button(self):
        page = self.client.get("/manage/classes-grades/")
        self.assertContains(page, 'href="/manage/classes-grades/grade/primary_3/"')
        self.assertContains(page, 'href="/manage/classes-grades/grade/creche/"')

    def test_renaming_a_grade_shows_everywhere(self):
        self.client.post("/manage/classes-grades/grade/primary_3/", {"label": "Basic 3", "level": "Level 3"})
        self.p3.refresh_from_db()
        self.assertEqual(self.p3.display_name, "Basic 3")
        self.assertEqual(User.objects.get(pk=self.ada.pk).level_display, "Basic 3")
        self.assertContains(self.client.get("/manage/classes-grades/"), "(changed)")
        self.ada.refresh_from_db()
        self.assertEqual(self.ada.level, "Level 3")                         # the lessons don't change

    def test_changing_its_level_moves_its_students_and_teachers(self):
        self.client.post("/manage/classes-grades/grade/primary_3/", {"label": "Grade 3", "level": "Level 4"})
        self.ada.refresh_from_db(), self.teacher.refresh_from_db()
        self.assertEqual((self.ada.level, self.ada.school_class), ("Level 4", self.p3))
        self.assertEqual(self.teacher.level, "Level 4")
        # Back to the default: they go back too, and nothing is stored.
        self.client.post("/manage/classes-grades/grade/primary_3/", {"action": "reset"})
        self.ada.refresh_from_db()
        self.assertEqual(self.ada.level, "Level 3")
        from .models import GradeSetting

        self.assertFalse(GradeSetting.objects.exists())

    def test_two_grades_cant_share_a_name(self):
        page = self.client.post("/manage/classes-grades/grade/primary_3/", {"label": "grade 7", "level": "Level 3"})
        self.assertContains(page, "Another grade is already called")

    def test_a_student_on_a_level_no_grade_follows_keeps_their_class(self):
        self.client.post("/manage/classes-grades/grade/primary_2/", {"label": "Grade 2", "level": "Level 3"})
        bo = person("bo@example.com", role="student", school=self.school, level="Level 2")   # Level 2 has no grade now
        self.assertIsNone(bo.school_class)
        self.ada.level = "Level 2"
        self.ada.save()
        self.ada.refresh_from_db()
        self.assertEqual(self.ada.school_class, self.p3)
