"""Excel import and first-login credential exports for school registration."""

import re
import secrets
from collections import OrderedDict
from datetime import date, datetime
from io import BytesIO

from django import forms
from django.core.exceptions import ValidationError
from django.core.validators import EmailValidator

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_SCHOOLS = 1000
MAX_ERRORS = 100

FIELDS = ("name", "address", "contact_person", "phone", "email", "relationship_status")
FIELD_LABELS = {
    "name": "School name",
    "address": "Location / address",
    "contact_person": "Contact person",
    "phone": "Phone number",
    "email": "Email",
    "relationship_status": "Status",
}
HEADER_ALIASES = {
    "name": {"schoolname", "name"},
    "address": {"locationaddress", "address", "location"},
    "contact_person": {"contactperson", "contactname", "contact"},
    "phone": {"phonenumber", "phone", "mobile", "telephone"},
    "email": {"email", "emailaddress"},
    "relationship_status": {"status", "relationshipstatus"},
}
MAX_LENGTHS = {"name": 255, "address": 255, "contact_person": 255, "phone": 120,
               "email": 254, "relationship_status": 30}
PASSWORD_WORDS = (
    "ant", "ape", "bee", "big", "box", "bud", "bus", "cap", "car", "cat",
    "cow", "cup", "day", "den", "dew", "dig", "dog", "dot", "ear", "egg",
    "elf", "fan", "fig", "fin", "fit", "fly", "fog", "fox", "fun", "gem",
    "gum", "hat", "hen", "hop", "hot", "hug", "ice", "ink", "jam", "jar",
    "jet", "job", "joy", "key", "kid", "kit", "leg", "log", "map", "mat",
    "mix", "mop", "mud", "mug", "net", "nut", "oak", "owl", "pan", "paw",
    "pea", "pen", "pet", "pie", "pig", "pin", "pot", "pup", "rag", "ram",
    "rat", "red", "rib", "rob", "rod", "row", "rub", "rug", "run", "sad",
    "sea", "sky", "sun", "tap", "tea", "ten", "tie", "tin", "toe", "top",
    "toy", "tub", "van", "web", "wet", "win", "yak", "yam", "yes", "zip", "zoo",
)


class SchoolWorkbookError(ValueError):
    """The uploaded file is not a readable school workbook."""


class SchoolWorkbookForm(forms.Form):
    workbook = forms.FileField(label="Excel workbook")

    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.fields["workbook"].widget.attrs.update({
            "class": "cr-file",
            "accept": ".xlsx,application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        })

    def clean_workbook(self):
        upload = self.cleaned_data["workbook"]
        if not upload.name.lower().endswith(".xlsx"):
            raise forms.ValidationError("Choose an Excel workbook ending in .xlsx.")
        if upload.size > MAX_FILE_BYTES:
            raise forms.ValidationError("The workbook must be 10 MB or smaller.")
        return upload


def _normal_header(value):
    return re.sub(r"[^a-z0-9]+", "", str(value or "").casefold())


def normalize_school_name(value):
    """A stable comparison key for duplicates in an upload and the database."""
    return " ".join(str(value or "").split()).casefold()


def _cell_value(cell, field):
    value = cell.value
    if value is None:
        return ""
    if cell.data_type == "f":
        raise ValueError("Formulas are not accepted; enter the value directly.")
    if isinstance(value, (datetime, date)):
        value = value.isoformat()
    if field == "phone" and isinstance(value, (int, float)) and not isinstance(value, bool):
        number = str(int(value)) if float(value).is_integer() else str(value)
        mask = cell.number_format or ""
        if re.fullmatch(r"0+", mask):
            number = number.zfill(len(mask))
        value = number
    return str(value).strip()


def _header_map(worksheet):
    for row_number, row in enumerate(worksheet.iter_rows(min_row=1, max_row=min(10, worksheet.max_row or 1)), start=1):
        found = {}
        for column, cell in enumerate(row):
            normalized = _normal_header(cell.value)
            for field, aliases in HEADER_ALIASES.items():
                if normalized in aliases:
                    found[field] = column
                    break
        if "name" in found:
            return row_number, found
    return None, {}


def parse_workbook(upload):
    """Return (unique records, row errors), reading all visible data sheets.

    Repeated school names across tabs are merged. Non-empty conflicting values
    are reported so a workbook never silently overwrites one row with another.
    """
    from openpyxl import load_workbook
    upload.seek(0)
    try:
        workbook = load_workbook(upload, read_only=True, data_only=False)
    except Exception as error:
        raise SchoolWorkbookError("This file could not be opened as an .xlsx workbook.") from error

    records = OrderedDict()
    errors = []
    sheet_count = 0
    source_rows = 0
    try:
        for worksheet in workbook.worksheets:
            if worksheet.sheet_state != "visible":
                continue
            header_row, columns = _header_map(worksheet)
            if header_row is None:
                continue
            sheet_count += 1
            if (worksheet.max_row or 0) > MAX_SCHOOLS + 10 or (worksheet.max_column or 0) > 100:
                errors.append(f"{worksheet.title}: keep each data sheet to {MAX_SCHOOLS} rows or fewer and 100 columns or fewer.")
                continue
            for row_number, row in enumerate(worksheet.iter_rows(min_row=header_row + 1), start=header_row + 1):
                values = {}
                row_has_data = False
                cell_errors = []
                for field, column in columns.items():
                    cell = row[column] if column < len(row) else None
                    if cell is None:
                        values[field] = ""
                        continue
                    try:
                        values[field] = _cell_value(cell, field)
                    except ValueError as error:
                        values[field] = ""
                        cell_errors.append(f"{FIELD_LABELS[field]}: {error}")
                    row_has_data = row_has_data or bool(values[field]) or bool(cell_errors)
                if not row_has_data:
                    continue
                source_rows += 1
                if source_rows > MAX_SCHOOLS:
                    errors.append(f"The workbook contains more than {MAX_SCHOOLS} school rows.")
                    break
                label = f"{worksheet.title}, row {row_number}"
                if cell_errors:
                    errors.extend(f"{label}: {message}" for message in cell_errors)
                    continue
                name = values.get("name", "")
                if not name:
                    errors.append(f"{label}: School name is required.")
                    continue
                for field in FIELDS:
                    value = values.get(field, "")
                    if len(value) > MAX_LENGTHS[field]:
                        errors.append(f"{label}: {FIELD_LABELS[field]} is longer than {MAX_LENGTHS[field]} characters.")
                status = values.get("relationship_status", "").casefold()
                if status and status not in {"new", "returning"}:
                    errors.append(f"{label}: Status must be New, Returning, or blank.")
                values["relationship_status"] = status
                email = values.get("email", "")
                if email:
                    try:
                        EmailValidator()(email)
                    except ValidationError:
                        errors.append(f"{label}: Email is not a valid email address.")
                    values["email"] = email.casefold()
                if errors and any(error.startswith(label + ":") for error in errors):
                    continue

                key = normalize_school_name(name)
                if key not in records:
                    values["name"] = " ".join(name.split())
                    values["source"] = label
                    records[key] = values
                    continue

                current = records[key]
                conflicts = []
                for field in FIELDS:
                    old_value, new_value = current.get(field, ""), values.get(field, "")
                    old_key = " ".join(old_value.split()).casefold()
                    new_key = " ".join(new_value.split()).casefold()
                    if old_value and new_value and old_key != new_key:
                        conflicts.append(FIELD_LABELS[field])
                    elif not old_value and new_value:
                        current[field] = new_value
                if conflicts:
                    errors.append(
                        f"{label}: Duplicate school name conflicts with {current['source']} in "
                        f"{', '.join(conflicts)}. Keep one consistent record."
                    )
            if source_rows > MAX_SCHOOLS:
                break
    finally:
        workbook.close()

    if sheet_count == 0:
        raise SchoolWorkbookError("No visible sheet has a School Name column in the first 10 rows.")
    if not records and not errors:
        errors.append("No school rows were found under the School Name heading.")
    return list(records.values()), errors[:MAX_ERRORS]


def template_file():
    """Return a clean Excel template with the supported school columns."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Schools"
    headings = ["School Name", "Location / Address", "Contact Person", "Phone Number", "Email", "Status"]
    widths = [34, 48, 30, 28, 36, 18]
    for column, (heading, width) in enumerate(zip(headings, widths), start=1):
        cell = worksheet.cell(row=1, column=column, value=heading)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="0B1A4A")
        cell.alignment = Alignment(vertical="center")
        worksheet.column_dimensions[cell.column_letter].width = width
    worksheet.row_dimensions[1].height = 28
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = "A1:F1001"
    status = DataValidation(type="list", formula1='"New,Returning"', allow_blank=True)
    status.promptTitle = "School status"
    status.prompt = "Choose New or Returning, or leave blank."
    status.errorTitle = "Choose a listed status"
    status.error = "Use New, Returning, or leave this cell blank."
    worksheet.add_data_validation(status)
    status.add("F2:F1001")
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def new_school_code(reserved_codes):
    """Generate a school code without one database lookup per school."""
    from apps.schools.utils import ALPHABET

    while True:
        code = "DM-" + "".join(secrets.choice(ALPHABET) for _ in range(6))
        if code not in reserved_codes:
            reserved_codes.add(code)
            return code


def new_school_admin_login(contact_person, school_name, reserved_usernames, reserved_emails):
    """Build a unique login username and short generated password.

    The school contact's name becomes the username. If the workbook has no
    contact name, use a school-based ``<school>.admin`` username instead.
    Existing and newly reserved identifiers are checked in memory so imports
    avoid a series of database uniqueness queries.
    """
    from django.contrib.auth import password_validation
    from django.utils.text import slugify

    from apps.accounts.models import INTERNAL_EMAIL_DOMAIN, User

    contact = " ".join(str(contact_person or "").split())
    if contact:
        parts = contact.split()
        first_name = parts[0][:150]
        last_name = " ".join(parts[1:])[:150]
        username_first, username_last = parts[0], " ".join(parts[1:])
        exported_name = contact
    else:
        first_name, last_name = "School", "Admin"
        username_first, username_last = school_name, "admin"
        exported_name = "School Admin"

    parts = [slugify(part) for part in (username_first, username_last)]
    parts = [part for part in parts if part]
    base = ".".join(parts).replace("-", "")[:34]
    base = re.sub(r"[^a-z0-9.]", "", base) or "student"
    suffix = 1
    while True:
        if suffix == 1:
            username = base
        else:
            suffix_text = str(suffix)
            username = f"{base[:40 - len(suffix_text)]}{suffix_text}"
        username_key = username.casefold()
        login_email = f"{username.lower()}@{INTERNAL_EMAIL_DOMAIN}"
        if username_key not in reserved_usernames and login_email.casefold() not in reserved_emails:
            reserved_usernames.add(username_key)
            reserved_emails.add(login_email.casefold())
            break
        suffix += 1

    # Bulk-created school accounts use a private internal email. The school's
    # supplied contact email remains on the School record; the admin signs in
    # with the username and password in the credential workbook.
    validation_user = User(
        email=login_email,
        username=username,
        first_name=first_name,
        last_name=last_name,
        role=User.Role.SCHOOL_ADMIN,
    )
    for _attempt in range(100):
        word = secrets.choice(PASSWORD_WORDS)
        digits = str(secrets.randbelow(90_000) + 10_000)
        password = f"{word.title()}{digits}"
        try:
            password_validation.validate_password(password, validation_user)
        except ValidationError:
            continue
        return {
            "contact_person": exported_name,
            "username": username,
            "email": login_email,
            "first_name": first_name,
            "last_name": last_name,
            "password": password,
        }
    raise RuntimeError("Could not generate a password accepted by the configured validators.")


def hash_school_admin_passwords(accounts, max_workers=4):
    """Hash credentials concurrently while retaining the configured hasher."""
    if not accounts:
        return []

    from concurrent.futures import ThreadPoolExecutor

    from django.contrib.auth.hashers import make_password

    worker_count = min(max_workers, len(accounts))
    with ThreadPoolExecutor(max_workers=worker_count) as pool:
        return list(pool.map(make_password, (account["password"] for account in accounts)))


def credentials_file(accounts):
    """Create a private Excel download containing newly created school logins."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "School logins"

    headings = ["School", "School code", "Contact person", "Username", "Password", "School contact email"]
    last_column = len(headings)
    worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_column)
    title = worksheet.cell(row=1, column=1, value="Diction Masters · School login details")
    title.font = Font(name="Aptos Display", size=16, bold=True, color="FFFFFF")
    title.fill = PatternFill("solid", fgColor="17183D")
    title.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 32

    worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_column)
    notice = worksheet.cell(
        row=2,
        column=1,
        value=(
            "Confidential: only new school accounts are listed. Sign in with Username and Password. "
            "A blank contact name uses a school-based admin username. The contact email is for reference. "
            "Current logins can be downloaded later from that school's people page. Keep this file secure."
        ),
    )
    notice.font = Font(name="Aptos", size=10, color="6A4D00", italic=True)
    notice.fill = PatternFill("solid", fgColor="FFF4C2")
    notice.alignment = Alignment(wrap_text=True, vertical="center")
    worksheet.row_dimensions[2].height = 42

    header_row = 4
    for column, heading in enumerate(headings, start=1):
        cell = worksheet.cell(row=header_row, column=column, value=heading)
        cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4D00D8")
        cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[header_row].height = 24

    widths = [34, 18, 30, 28, 28, 36]
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[worksheet.cell(row=header_row, column=column).column_letter].width = width

    bottom = Side(style="hair", color="DCE0EF")
    for row_number, account in enumerate(accounts, start=header_row + 1):
        values = [
            account["school"].name,
            account["school"].code,
            account["contact_person"],
            account["username"],
            account["password"],
            account["school"].email,
        ]
        for column, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_number, column=column)
            cell.value = str(value or "")
            # Keep arbitrary workbook content and passwords as literal text.
            cell.data_type = "s"
            cell.number_format = "@"
            cell.alignment = Alignment(vertical="top", wrap_text=column in (1, 3, 6))
            cell.border = Border(bottom=bottom)
            if column == 5:
                cell.font = Font(name="Aptos Mono", color="17183D")
        if row_number % 2:
            for cell in worksheet[row_number][:last_column]:
                cell.fill = PatternFill("solid", fgColor="F5F6FC")

    worksheet.freeze_panes = f"A{header_row + 1}"
    if accounts:
        worksheet.auto_filter.ref = f"A{header_row}:F{header_row + len(accounts)}"
    worksheet.sheet_view.showGridLines = False

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
